"""Run: python -m pytest -q. No real OpenAI billing during tests."""
import base64
import pytest
from fastapi.testclient import TestClient
from app import app
from data import METHODS
import engine

client=TestClient(app)

@pytest.fixture(scope='module')
def demo_index():
    result=client.post('/api/index',json={'mode':'offline','use_sample':True,'documents':[]})
    assert result.status_code==200, result.text
    return result.json()['index_id']


def test_bootstrap_matches_13_papers():
    obj=client.get('/api/bootstrap').json()
    assert len(obj['methods'])==14 # baseline + 13 research methods
    assert len(obj['documents'])>=8
    assert all(x['example'] and x['source'].startswith('https://') for x in obj['methods'])


def test_13_method_end_to_end_offline(demo_index):
    for method in METHODS:
        response=client.post('/api/ask',json={'index_id':demo_index,'method':method['id'],'question':method['example']})
        assert response.status_code==200, (method['id'],response.text)
        data=response.json()
        assert data['method']==method['id']
        assert data['evidence'], method['id']
        assert data['answer'].startswith('OFFLINE')
        assert len(data['trace'])>=3
        assert data['metrics']['mode']=='offline'


def test_compare_and_provenance(demo_index):
    response=client.post('/api/compare',json={'index_id':demo_index,'methods':['baseline','hipporag','graphrag'],'question':'TrafficStream과 Atlas 관계를 설명해줘'})
    assert response.status_code==200,response.text
    results=response.json()['results']
    assert len(results)==3
    assert all(e['id'].startswith(('D','A','C')) for r in results for e in r['evidence'])


def test_custom_documents_have_generic_offline_graph():
    response=client.post('/api/index',json={'mode':'offline','use_sample':False,'documents':[{'title':'테스트','text':'Alpha uses Beta. Beta supports Gamma.'},{'title':'테스트2','text':'Gamma and Delta cooperate.'}]})
    assert response.status_code==200,response.text
    assert response.json()['stats']['documents']==2
    assert response.json()['stats']['relations']>=1


def test_image_offline_explicitly_not_analyzed():
    payload='data:image/png;base64,'+base64.b64encode(b'fake small png').decode()
    response=client.post('/api/index',json={'mode':'offline','use_sample':True,'image_data_url':payload})
    assert response.status_code==200
    assert response.json()['stats']['vision'] is False


def test_live_requires_server_key(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    response=client.post('/api/index',json={'mode':'live','use_sample':True})
    assert response.status_code==400
    assert 'OPENAI_API_KEY' in response.json()['detail']


def test_real_call_path_mocked(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','sk-TEST-NOT-REAL')
    calls={'extract':0,'embeds':0,'answer':0,'vision':0}
    def extract(docs):
        calls['extract']+=1
        return ([{'source':'Atlas','relation':'사용','target':'Beacon','doc_id':'D1'}],
                [{'id':'A1','doc_id':'D1','text':docs[0]['text']}])
    def embs(items):
        calls['embeds']+=1
        return [[float(len(text)%11),2.,5.] for text in items]
    def answer(result,query):
        calls['answer']+=1
        return f'테스트용 실제 호출 경로 스텁: {query} [D1]'
    def vision(url):
        calls['vision']+=1
        return '도식에는 Atlas에서 Beacon으로 연결된 화살표가 있다.'
    monkeypatch.setattr(engine,'extract_live',extract)
    monkeypatch.setattr(engine,'embed_many',embs)
    monkeypatch.setattr(engine,'answer_live',answer)
    monkeypatch.setattr(engine,'caption_image',vision)
    image='data:image/png;base64,'+base64.b64encode(b'fake bytes').decode()
    response=client.post('/api/index',json={'mode':'live','use_sample':True,'image_data_url':image})
    assert response.status_code==200,response.text
    assert response.json()['stats']['vision'] is True
    idx=response.json()['index_id']
    for m in METHODS:
        res=client.post('/api/ask',json={'index_id':idx,'question':m['example'],'method':m['id']})
        assert res.status_code==200,(m['id'],res.text)
        assert res.json()['metrics']['answer_kind']=='generated'
    assert calls['extract']==1 and calls['vision']==1 and calls['embeds']>=14 and calls['answer']==14


def test_validation(demo_index):
    assert client.post('/api/ask',json={'index_id':demo_index,'method':'not-real','question':'test'}).status_code==400
    assert client.post('/api/ask',json={'index_id':'garbage','method':'raptor','question':'test'}).status_code==404
    assert client.post('/api/compare',json={'index_id':demo_index,'methods':['baseline','baseline'],'question':'test'}).status_code==400


def test_home_static_assets_and_security():
    home=client.get('/')
    assert home.status_code==200
    assert 'GRAPHATLAS' in home.text
    assert 'OPENAI_API_KEY' not in home.text
    assert client.get('/static/app.js').status_code==200
    assert client.get('/static/styles.css').status_code==200
    assert client.get('/static/assets/knowledge-universe.png').status_code==200


def test_research_sources_architecture():
    r=client.get('/api/research')
    assert r.status_code==200
    data=r.json()
    assert len(data)==13
    assert all(len(x['index'])>=5 and len(x['query'])>=5 for x in data)
    assert all(x['paper'].startswith('https://arxiv.org/') and x['repo'].startswith('https://github.com/') for x in data)


def test_offline_sse_stream_retrieval_delta_done(demo_index):
    response=client.post('/api/ask/stream',json={'index_id':demo_index,'method':'hipporag','question':'Orion이 Beacon에 전달하는 것은?'})
    assert response.status_code==200
    assert response.headers['content-type'].startswith('text/event-stream')
    txt=response.text
    assert 'event: retrieval' in txt and 'event: delta' in txt and 'event: done' in txt
    assert 'OFFLINE' in txt
    assert 'event: error' not in txt


def test_stream_validation(demo_index):
    res=client.post('/api/ask/stream',json={'index_id':demo_index,'method':'bogus','question':'hello'})
    assert res.status_code==400


def test_presentation_download_and_presenter_notes():
    from pptx import Presentation
    from io import BytesIO
    response=client.get('/api/download/slides')
    assert response.status_code==200
    assert response.content[:2]==b'PK'
    deck=Presentation(BytesIO(response.content))
    assert len(deck.slides)==60
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)
    assert client.get('/api/health').json()['ok']

def test_gpt6_luna_compatible_params(monkeypatch):
    monkeypatch.setenv('OPENAI_CHAT_MODEL','gpt-6-luna')
    settings=engine.chat_params()
    assert settings['model']=='gpt-6-luna'
    assert settings['reasoning_effort']=='none'
    assert 'temperature' not in settings

def test_public_deploy_basicauth(monkeypatch):
    import base64
    monkeypatch.setenv('DEMO_PASSWORD','local-test-only')
    monkeypatch.setenv('RENDER','true')
    assert client.get('/').status_code==401
    assert client.get('/api/health').status_code==200
    cred=base64.b64encode(b'seminar:local-test-only').decode()
    assert client.get('/api/bootstrap',headers={'Authorization':'Basic '+cred}).status_code==200
